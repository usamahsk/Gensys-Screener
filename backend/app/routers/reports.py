"""Dashboard reports: downloadable Excel/CSV exports.

Report chooser (frontend /reports page):
  - orders — every order + its items, provider/telephone ids included
  - calls  — every call log with provider/telephone ids, costs, summaries
  - full   — everything above plus a Summary sheet (KPIs, daily series,
             disposition + provider breakdowns) — the informative default

Formats: xlsx (multi-sheet workbook) or csv (single-sheet).
Optional ?start_date= / ?end_date= (ISO date or datetime) filter both tables.
"""

from datetime import datetime, timezone
from io import BytesIO, StringIO
from typing import Optional

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.db import get_db
from app.models import Call, Order

router = APIRouter(prefix="/api/reports", tags=["reports"])

VALID_REPORTS = {"orders", "calls", "full"}


def _parse_dt(v: Optional[str]) -> Optional[datetime]:
    if not v or not v.strip():
        return None
    s = v.strip()
    try:
        dt = datetime.fromisoformat(s.replace("Z", "+00:00"))
    except ValueError:
        try:
            dt = datetime.strptime(s[:10], "%Y-%m-%d")
        except ValueError:
            raise HTTPException(400, f"Bad date {v!r} — use YYYY-MM-DD")
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=timezone.utc)
    return dt


def _style_header(ws, ncols: int) -> None:
    from openpyxl.styles import Font, PatternFill

    fill = PatternFill("solid", fgColor="0044BA")
    font = Font(bold=True, color="FFFFFF", size=11)
    for col in range(1, ncols + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = fill
        cell.font = font
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def _autosize(ws, widths: list[int]) -> None:
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w


@router.get("/download")
async def download_report(
    report: str = Query(default="full"),
    format: str = Query(default="xlsx"),
    start_date: Optional[str] = None,
    end_date: Optional[str] = None,
    db: AsyncSession = Depends(get_db),
):
    report = (report or "full").lower()
    fmt = (format or "xlsx").lower()
    if report not in VALID_REPORTS:
        raise HTTPException(400, f"Unknown report {report!r} — use orders, calls or full")
    if fmt not in ("xlsx", "csv"):
        raise HTTPException(400, f"Unknown format {fmt!r} — use xlsx or csv")
    if fmt == "csv" and report == "full":
        raise HTTPException(400, "CSV supports report=orders or report=calls — use xlsx for full")

    start = _parse_dt(start_date)
    end = _parse_dt(end_date)

    oq = select(Order).options(selectinload(Order.items)).order_by(Order.created_at.desc())
    cq = select(Call).order_by(Call.started_at.desc())
    if start:
        oq = oq.where(Order.created_at >= start)
        cq = cq.where(Call.started_at >= start)
    if end:
        oq = oq.where(Order.created_at <= end)
        cq = cq.where(Call.started_at <= end)
    orders = (await db.execute(oq)).scalars().all()
    calls = (await db.execute(cq)).scalars().all()

    stamp = datetime.now(timezone.utc).strftime("%Y%m%d-%H%M")
    if fmt == "csv":
        return _csv_response(report, orders, calls, stamp)
    return _xlsx_response(report, orders, calls, stamp, start_date, end_date)


# ---------- CSV ----------

def _csv_response(report: str, orders, calls, stamp: str) -> StreamingResponse:
    import csv as _csv

    buf = StringIO()
    if report == "orders":
        rows = _order_rows(orders)
        w = _csv.DictWriter(buf, fieldnames=list(rows[0].keys()) if rows else ["note"])
        w.writeheader()
        w.writerows(rows or [{"note": "no orders in range"}])
        filename = f"orders-report-{stamp}.csv"
    else:
        rows = _call_rows(calls)
        w = _csv.DictWriter(buf, fieldnames=list(rows[0].keys()) if rows else ["note"])
        w.writeheader()
        w.writerows(rows or [{"note": "no calls in range"}])
        filename = f"calls-report-{stamp}.csv"
    data = buf.getvalue().encode("utf-8")
    return StreamingResponse(
        iter([data]), media_type="text/csv",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


# ---------- XLSX ----------

def _xlsx_response(report: str, orders, calls, stamp: str,
                   start_date: Optional[str], end_date: Optional[str]) -> StreamingResponse:
    from openpyxl import Workbook

    wb = Workbook()
    if report in ("orders", "full"):
        ws = wb.active
        ws.title = "Orders"
        _fill_orders_sheet(ws, orders)
    if report in ("calls", "full"):
        ws = wb.create_sheet("Calls") if report == "full" else wb.active
        if report == "calls":
            ws.title = "Calls"
        _fill_calls_sheet(ws, calls)
    if report == "full":
        _fill_items_sheet(wb.create_sheet("Order Items"), orders)
        _fill_summary_sheet(wb.create_sheet("Summary"), orders, calls,
                            start_date, end_date)

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = f"{report}-report-{stamp}.xlsx"
    return StreamingResponse(
        iter([buf.getvalue()]),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


def _order_rows(orders) -> list[dict]:
    rows = []
    for o in orders:
        items = "; ".join(f"{i.quantity}x {i.name}" + (f" ({i.size})" if i.size else "")
                          for i in o.items) or "—"
        rows.append({
            "Order #": o.order_number,
            "Placed at": o.created_at.isoformat(timespec="seconds") if o.created_at else "",
            "Customer": o.customer_name or "",
            "Phone": o.customer_phone or "",
            "Type": o.order_type or "",
            "Table / Pickup": o.table_number or o.pickup_time or "",
            "Status": o.status or "",
            "Items count": len(o.items),
            "Items": items,
            "Total": float(o.total or 0),
            "Payment": o.payment_method or "",
            "Provider": o.provider or "dograh",
            "Provider ID": o.provider_call_id or "",
            "Call ID": o.call_id or "",
            "Source": o.source or "",
        })
    return rows


def _call_rows(calls) -> list[dict]:
    rows = []
    for c in calls:
        rows.append({
            "Call ID": c.call_id or "",
            "Started at": c.started_at.isoformat(timespec="seconds") if c.started_at else "",
            "Provider": c.provider or "dograh",
            "Provider ID": c.provider_call_id or "",
            "Genesys conversation ID": c.genesys_conversation_id or "",
            "Run ID": c.run_id or "",
            "Workflow": c.workflow_name or c.workflow_id or "",
            "Call type": c.call_type or "",
            "Disposition": c.disposition or "",
            "Customer": c.customer_name or c.first_name or "",
            "Phone": c.customer_phone or "",
            "Duration (s)": c.duration_seconds or 0,
            "Orders": ", ".join(c.order_numbers or ""),
            "Sentiment": c.sentiment or "",
            "Cost": float(c.cost) if c.cost is not None else "",
            "Summary": (c.summary or "")[:500],
        })
    return rows


def _fill_orders_sheet(ws, orders) -> None:
    rows = _order_rows(orders)
    headers = list(rows[0].keys()) if rows else ["Note"]
    ws.append(headers)
    for r in rows:
        ws.append(list(r.values()))
    if not rows:
        ws.append(["No orders in range"])
    _style_header(ws, len(headers))
    _autosize(ws, [14, 20, 18, 16, 10, 14, 12, 11, 50, 10, 10, 14, 28, 20, 12])


def _fill_calls_sheet(ws, calls) -> None:
    rows = _call_rows(calls)
    headers = list(rows[0].keys()) if rows else ["Note"]
    ws.append(headers)
    for r in rows:
        ws.append(list(r.values()))
    if not rows:
        ws.append(["No calls in range"])
    _style_header(ws, len(headers))
    _autosize(ws, [22, 20, 12, 28, 24, 10, 16, 12, 12, 18, 16, 11, 16, 10, 8, 60])


def _fill_items_sheet(ws, orders) -> None:
    headers = ["Order #", "Item", "Size", "Qty", "Unit price", "Line total", "SKU", "Notes"]
    ws.append(headers)
    for o in orders:
        for i in o.items:
            ws.append([o.order_number, i.name, i.size or "", i.quantity,
                       float(i.unit_price or 0),
                       round(float(i.unit_price or 0) * i.quantity, 2),
                       i.sku or "", i.notes or ""])
    if ws.max_row == 1:
        ws.append(["—", "No items in range", "", "", "", "", "", ""])
    _style_header(ws, len(headers))
    _autosize(ws, [14, 34, 12, 8, 12, 12, 18, 30])


def _fill_summary_sheet(ws, orders, calls,
                        start_date: Optional[str], end_date: Optional[str]) -> None:
    from openpyxl.styles import Font

    title = Font(bold=True, size=13, color="000E49")
    bold = Font(bold=True)

    ws.append(["Voice commerce report — Summary"])
    ws["A1"].font = title
    ws.append(["Period", f"{start_date or 'beginning'} → {end_date or 'now'}"])
    ws.append(["Generated (UTC)", datetime.now(timezone.utc).isoformat(timespec="seconds")])
    ws.append([])

    revenue = sum(float(o.total or 0) for o in orders)
    total_dur = sum(c.duration_seconds or 0 for c in calls)
    ws.append(["KPI", "Value"])
    ws.append(["Total calls", len(calls)])
    ws.append(["AI resolved (completed)", sum(1 for c in calls if c.disposition == "completed")])
    ws.append(["Transferred", sum(1 for c in calls if c.disposition == "transferred")])
    ws.append(["Failed / abandoned",
               sum(1 for c in calls if c.disposition in ("failed", "abandoned"))])
    ws.append(["Orders created", len(orders)])
    ws.append(["Revenue", round(revenue, 2)])
    ws.append(["Avg order value", round(revenue / len(orders), 2) if orders else 0])
    ws.append(["Avg call duration (s)",
               round(total_dur / len(calls)) if calls else 0])
    for row in ws.iter_rows(min_row=5, max_row=ws.max_row, max_col=1):
        row[0].font = bold
    ws.append([])

    ws.append(["Calls by disposition", "Count"])
    disp: dict[str, int] = {}
    for c in calls:
        disp[c.disposition or "unknown"] = disp.get(c.disposition or "unknown", 0) + 1
    for k in sorted(disp):
        ws.append([k, disp[k]])
    ws.append([])

    ws.append(["Calls by provider", "Count"])
    prov: dict[str, int] = {}
    for c in calls:
        prov[c.provider or "dograh"] = prov.get(c.provider or "dograh", 0) + 1
    for k in sorted(prov):
        ws.append([k, prov[k]])
    ws.append([])

    ws.append(["Orders by status", "Count", "Revenue"])
    ostat: dict[str, list] = {}
    for o in orders:
        e = ostat.setdefault(o.status or "unknown", [0, 0.0])
        e[0] += 1
        e[1] += float(o.total or 0)
    for k in sorted(ostat):
        ws.append([k, ostat[k][0], round(ostat[k][1], 2)])
    ws.append([])

    ws.append(["Date", "Calls", "Orders", "Revenue"])
    daily: dict[str, list] = {}
    for c in calls:
        if c.started_at:
            e = daily.setdefault(c.started_at.date().isoformat(), [0, 0, 0.0])
            e[0] += 1
    for o in orders:
        if o.created_at:
            e = daily.setdefault(o.created_at.date().isoformat(), [0, 0, 0.0])
            e[1] += 1
            e[2] += float(o.total or 0)
    for d in sorted(daily):
        ws.append([d, daily[d][0], daily[d][1], round(daily[d][2], 2)])

    _autosize(ws, [28, 18, 18])
